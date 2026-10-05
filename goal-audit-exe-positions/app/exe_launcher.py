# -*- coding: utf-8 -*-
import csv
import re
import sys
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox
import analyze_goals as core

POS_ALIASES = ['должность', 'наименование должности']

def clean(value):
    return '' if value is None else re.sub(r'\s+', ' ', str(value).replace('\u00a0', ' ')).strip()

def normalize_position(value):
    return clean(value).replace('ё', 'е').casefold()

def read_raw(path):
    path = Path(path)
    if path.suffix.lower() == '.csv':
        for encoding in ('utf-8-sig', 'cp1251'):
            try:
                with path.open(encoding=encoding, newline='') as stream:
                    sample = stream.read(4096)
                    stream.seek(0)
                    try:
                        dialect = csv.Sniffer().sniff(sample, delimiters=';,\t,')
                    except csv.Error:
                        dialect = csv.excel
                        dialect.delimiter = ';'
                    return list(csv.reader(stream, dialect))
            except UnicodeDecodeError:
                continue
        raise ValueError('Не удалось определить кодировку CSV-файла.')
    if path.suffix.lower() not in ('.xlsx', '.xlsm'):
        raise ValueError('Допустимы только файлы XLSX, XLSM или CSV.')
    from openpyxl import load_workbook
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    workbook.close()
    return rows

def find_position_column(rows):
    best = None
    for row_index, row in enumerate(rows[:30]):
        values = [clean(value).casefold() for value in row]
        for alias in POS_ALIASES:
            for column_index, value in enumerate(values):
                if value == alias:
                    return row_index, column_index
        for alias in POS_ALIASES:
            for column_index, value in enumerate(values):
                if value and alias in value:
                    best = best or (row_index, column_index)
    if best:
        return best
    raise ValueError('В выбранном файле не найден столбец «Должность».')

def read_positions(path):
    rows = read_raw(path)
    if not rows:
        raise ValueError('Выбранный файл со списком должностей пуст.')
    header_row, column = find_position_column(rows)
    positions = set()
    for row in rows[header_row + 1:]:
        value = clean(row[column] if column < len(row) else '')
        if value:
            positions.add(normalize_position(value))
    if not positions:
        raise ValueError('Под столбцом «Должность» не найдено ни одной заполненной должности.')
    return positions

def choose_file(root, title, message):
    messagebox.showinfo(title, message, parent=root)
    selected = filedialog.askopenfilename(
        parent=root,
        title=title,
        filetypes=[('Таблицы', '*.xlsx *.xlsm *.csv'), ('Все файлы', '*.*')]
    )
    if not selected:
        raise RuntimeError('Выбор файла отменён.')
    return selected

def progress(done, total):
    width = 34
    filled = int(width * done / max(total, 1))
    bar = '#' * filled + '-' * (width - filled)
    percent = int(100 * done / max(total, 1))
    print(f'\rОбработка целей: [{bar}] {done}/{total} ({percent:3d}%)', end='', flush=True)
    if done == total:
        print()

def main():
    root = tk.Tk()
    root.withdraw()
    root.attributes('-topmost', True)
    try:
        positions_file = choose_file(
            root,
            'Шаг 1 из 2 — список должностей',
            'Выберите список должностей для анализа по каскадным метрикам'
        )
        allowed_positions = read_positions(positions_file)
        print(f'Выбрано уникальных должностей: {len(allowed_positions)}')

        goals_file = choose_file(
            root,
            'Шаг 2 из 2 — файл с целями',
            'Выберите входной файл с целями сотрудников. Обязательные столбцы: «Должность» и «Наименование цели».'
        )
        all_rows = core.read_rows(goals_file)
        if not all_rows:
            raise ValueError('Не удалось прочитать ни одной строки с целями.')
        selected_rows = [row for row in all_rows if normalize_position(row['pos']) in allowed_positions]
        excluded_rows = [row for row in all_rows if normalize_position(row['pos']) not in allowed_positions]
        print(f'Целей до отбора: {len(all_rows)}; после отбора: {len(selected_rows)}; исключено: {len(excluded_rows)}')

        result, statistics = core.analyze(selected_rows, progress)
        output = filedialog.asksaveasfilename(
            parent=root,
            title='Сохранить итоговый отчёт',
            defaultextension='.xlsx',
            initialfile='Аудит_целей_по_каскаду_метрик.xlsx',
            filetypes=[('Excel', '*.xlsx')]
        )
        if not output:
            raise RuntimeError('Сохранение отчёта отменено.')
        core.write_xlsx(result, statistics, output, excluded_rows=excluded_rows)
        messagebox.showinfo(
            'Готово',
            f'Проверено целей: {len(result)}\nНе вошли в оценку: {len(excluded_rows)}\n\nФайл сохранён:\n{output}',
            parent=root,
        )
        return 0
    except Exception as error:
        messagebox.showerror('Ошибка', str(error), parent=root)
        print('ОШИБКА:', error, file=sys.stderr)
        return 1
    finally:
        root.destroy()

if __name__ == '__main__':
    sys.exit(main())
