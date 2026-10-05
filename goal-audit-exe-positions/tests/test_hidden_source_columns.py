import importlib.util
from datetime import datetime
from pathlib import Path
import tempfile
import unittest
import sys
from unittest.mock import MagicMock, patch

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter


SCRIPT = Path(__file__).resolve().parents[1] / "app" / "analyze_goals.py"
SPEC = importlib.util.spec_from_file_location("analyze_goals", SCRIPT)
core = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(core)


class HiddenSourceColumnsTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.headers = [
            "ФИО", "Должность", "Наименование цели", "Описание цели",
            "Дополнительное поле", "Дата", "Число", "Логическое", None,
            "Повтор", "Повтор",
        ]
        self.values = [
            ["Иван", "Инженер", "Провести проверки", "Не менее 10 шт",
             "  исходный текст  ", datetime(2026, 10, 5), 12.5, True, None, "А", "Б"],
            ["Пётр", "Оператор", "Провести обход", "Не менее 5 шт",
             "вторая строка", None, 0, False, None, "В", "Г"],
        ]
        wb = Workbook()
        ws = wb.active
        ws.append(["Служебная строка"])
        ws.append(self.headers)
        for row in self.values:
            ws.append(row)
        # Такая строка и раньше исключалась из анализа.
        ws.append([None, None, None, None, "Не цель"])
        self.input = self.root / "input.xlsx"
        wb.save(self.input)
        wb.close()

    def export(self, rows, filename="result.xlsx", excluded_rows=None):
        result, statistics = core.analyze(rows)
        path = self.root / filename
        core.write_xlsx(result, statistics, path, excluded_rows=excluded_rows)
        wb = load_workbook(path)
        self.addCleanup(wb.close)
        return result, wb

    def test_all_source_columns_values_and_types_are_preserved_and_hidden(self):
        rows = core.read_rows(str(self.input))
        self.assertEqual([r["row"] for r in rows], [3, 4])
        result, wb = self.export(rows)
        ws = wb.active
        self.assertEqual(ws.max_column, 15 + len(self.headers))
        self.assertEqual(ws.max_row, 3)
        self.assertEqual([ws.cell(1, i).value for i in range(16, ws.max_column + 1)],
                         self.headers)
        for row_number, expected in enumerate(self.values, start=2):
            self.assertEqual([ws.cell(row_number, i).value
                              for i in range(16, ws.max_column + 1)], expected)
        for i in range(16, ws.max_column + 1):
            self.assertTrue(ws.column_dimensions[get_column_letter(i)].hidden)
        for i in range(1, 16):
            self.assertFalse(ws.column_dimensions[get_column_letter(i)].hidden)
        self.assertEqual(ws.auto_filter.ref, "A1:O3")
        self.assertEqual(ws.freeze_panes, "C2")
        self.assertEqual(result[0]["source_values"][4], "  исходный текст  ")
        self.assertIsInstance(ws.cell(2, 21).value, datetime)
        self.assertEqual(ws.cell(2, 23).data_type, "b")

    def test_filtering_keeps_the_correct_original_row(self):
        rows = core.read_rows(str(self.input), header_row=2)
        _, wb = self.export([r for r in rows if r["pos"] == "Оператор"])
        ws = wb.active
        self.assertEqual(ws.cell(2, 1).value, 4)
        self.assertEqual([ws.cell(2, i).value for i in range(16, ws.max_column + 1)],
                         self.values[1])

    def test_analysis_and_visible_report_are_unchanged(self):
        rows = core.read_rows(str(self.input))
        plain = [{k: v for k, v in r.items() if not k.startswith("source_")} for r in rows]
        result, wb = self.export(rows)
        baseline, old_wb = self.export(plain, "baseline.xlsx")
        strip = lambda r: {k: v for k, v in r.items() if not k.startswith("source_")}
        self.assertEqual([strip(r) for r in result], [strip(r) for r in baseline])
        for row in wb.active.iter_rows(max_col=15):
            for cell in row:
                other = old_wb.active.cell(cell.row, cell.column)
                self.assertEqual(cell.value, other.value)
                self.assertEqual(cell._style, other._style)
        self.assertEqual(list(wb["Сводка"].values), list(old_wb["Сводка"].values))
        self.assertEqual(old_wb.active.max_column, 15)

    def test_csv_ragged_rows_and_formula_like_text(self):
        path = self.root / "input.csv"
        path.write_text(
            "Должность;Наименование цели;Дополнительное поле\n"
            "Инженер;Провести проверки;=1+1;лишнее поле\n"
            "Оператор;Провести обход\n",
            encoding="utf-8-sig",
        )
        rows = core.read_rows(str(path))
        _, wb = self.export(rows)
        ws = wb.active
        self.assertEqual(ws.max_column, 19)
        self.assertEqual(ws.cell(1, 19).value, None)
        self.assertTrue(ws.column_dimensions["S"].hidden)
        self.assertEqual(ws.cell(2, 18).value, "=1+1")
        self.assertEqual(ws.cell(2, 18).data_type, "s")
        self.assertEqual(ws.cell(2, 19).value, "лишнее поле")
        self.assertEqual(ws.cell(3, 18).value, None)

    def test_visible_department_keeps_full_path_without_truncation(self):
        full_path = (
            "Компания/Завод/Производство № 1/"
            "Цех с длинным наименованием для проверки сохранения полного пути/"
            "Участок № 5"
        )
        self.assertGreater(len(full_path), 70)
        wb = load_workbook(self.input)
        ws = wb.active
        ws.cell(2, 12, "Подразделение")
        ws.cell(3, 12, full_path)
        ws.cell(4, 12, "Производство № 2")
        wb.save(self.input)
        wb.close()
        rows = core.read_rows(str(self.input), header_row=2)
        result, report = self.export(rows)
        self.assertEqual(result[0]["unit"], full_path)
        self.assertEqual(report.active.cell(2, 5).value, full_path)
        self.assertEqual(report.active.cell(3, 5).value, "Производство № 2")
        self.assertEqual(report.active.cell(2, 27).value, full_path)
        self.assertTrue(report.active.column_dimensions["AA"].hidden)

    def test_excluded_sheet_preserves_all_fields_order_and_duplicates(self):
        rows = core.read_rows(str(self.input))
        excluded = [rows[0], dict(rows[0])]
        _, report = self.export([rows[1]], excluded_rows=excluded)
        sheet = report["Не вошли в оценку"]
        self.assertEqual(list(sheet.values),
                         [tuple(self.headers), tuple(self.values[0]), tuple(self.values[0])])
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet.auto_filter.ref, "A1:K3")
        for i in range(1, len(self.headers) + 1):
            self.assertFalse(sheet.column_dimensions[get_column_letter(i)].hidden)
        self.assertEqual(report.active.max_row, 2)
        self.assertEqual(report.active.cell(2, 3).value, "Пётр")
        self.assertEqual(report["Сводка"].cell(2, 2).value, 1)
        self.assertEqual(report["Сводка"].cell(3, 2).value, 1)

    def test_no_exclusions_produces_header_only_sheet(self):
        _, report = self.export(core.read_rows(str(self.input)))
        self.assertEqual(list(report["Не вошли в оценку"].values), [tuple(self.headers)])

    def test_all_goals_excluded_produces_report_with_zero_summary(self):
        rows = core.read_rows(str(self.input))
        result, report = self.export([], excluded_rows=rows)
        self.assertEqual(result, [])
        self.assertEqual(report.active.max_row, 1)
        self.assertEqual(report["Сводка"].cell(2, 2).value, 0)
        self.assertEqual(report["Сводка"].cell(3, 2).value, 0)
        self.assertEqual(list(report["Не вошли в оценку"].values),
                         [tuple(self.headers)] + [tuple(row) for row in self.values])

    def test_launcher_passes_excluded_rows_to_export(self):
        launcher_path = SCRIPT.parent / "exe_launcher.py"
        spec = importlib.util.spec_from_file_location("exe_launcher", launcher_path)
        launcher = importlib.util.module_from_spec(spec)
        with patch.dict(sys.modules, {"analyze_goals": core}):
            spec.loader.exec_module(launcher)
        for allowed in ({"оператор"}, {"несуществующая должность"}):
            with self.subTest(allowed=allowed):
                output = self.root / "launcher.xlsx"
                with (
                    patch.object(launcher.tk, "Tk", return_value=MagicMock()),
                    patch.object(launcher, "choose_file", side_effect=["positions.xlsx", str(self.input)]),
                    patch.object(launcher, "read_positions", return_value=allowed),
                    patch.object(launcher.filedialog, "asksaveasfilename", return_value=str(output)),
                    patch.object(launcher.messagebox, "showinfo"),
                    patch.object(launcher.messagebox, "showerror") as showerror,
                ):
                    self.assertEqual(launcher.main(), 0)
                    showerror.assert_not_called()
                report = load_workbook(output)
                try:
                    expected_excluded = 1 if "оператор" in allowed else 2
                    self.assertEqual(report["Не вошли в оценку"].max_row, expected_excluded + 1)
                    self.assertEqual(report["Сводка"].cell(2, 2).value, 2 - expected_excluded)
                finally:
                    report.close()


if __name__ == "__main__":
    unittest.main()
